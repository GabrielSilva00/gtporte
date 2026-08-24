# -*- coding: utf-8 -*-
"""
Gera docs/relatorio-automacao-gtporte.docx a partir do Markdown ao lado.

Escrito porque o relatório precisa ser entregue em .docx, mas a fonte de
verdade deve continuar versionável e revisável em texto puro. Rodar de
novo regenera o documento inteiro — não edite o .docx à mão, edite o .md.

Uso:
    pip install python-docx
    python scripts/gerar-relatorio.py
"""
import re
import sys
from pathlib import Path

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Cm
except ImportError:
    sys.exit("python-docx não instalado. Rode: pip install python-docx")

RAIZ = Path(__file__).resolve().parent.parent
FONTE = RAIZ / "docs" / "relatorio-automacao-gtporte.md"
DESTINO = RAIZ / "docs" / "relatorio-automacao-gtporte.docx"

# Verde do tema do sistema, para os títulos ficarem coerentes com o produto.
PRIMARIA = RGBColor(0x1F, 0x3A, 0x2E)
CODIGO_BG = RGBColor(0x6B, 0x75, 0x70)


def configurar_estilos(d):
    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for nivel, tamanho in ((1, 18), (2, 14), (3, 12)):
        estilo = d.styles[f"Heading {nivel}"]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamanho)
        estilo.font.color.rgb = PRIMARIA
        estilo.font.bold = True
        estilo.paragraph_format.space_before = Pt(14 if nivel == 1 else 10)
        estilo.paragraph_format.space_after = Pt(4)


def escrever_inline(paragrafo, texto):
    """Interpreta **negrito**, *itálico* e `código` de uma linha."""
    for parte in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            paragrafo.add_run(parte[2:-2]).bold = True
        elif parte.startswith("`") and parte.endswith("`"):
            run = paragrafo.add_run(parte[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = CODIGO_BG
        elif parte.startswith("*") and parte.endswith("*"):
            paragrafo.add_run(parte[1:-1]).italic = True
        else:
            paragrafo.add_run(parte)


def adicionar_tabela(d, linhas):
    """Converte um bloco de tabela Markdown em tabela do Word."""
    celulas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas]
    cabecalho, corpo = celulas[0], celulas[2:]  # linha 1 é o separador

    tabela = d.add_table(rows=1, cols=len(cabecalho))
    tabela.style = "Light Grid Accent 1"
    tabela.autofit = True

    for i, titulo in enumerate(cabecalho):
        cel = tabela.rows[0].cells[i]
        cel.text = ""
        escrever_inline(cel.paragraphs[0], titulo)
        for run in cel.paragraphs[0].runs:
            run.bold = True

    for linha in corpo:
        cels = tabela.add_row().cells
        for i, valor in enumerate(linha[: len(cabecalho)]):
            cels[i].text = ""
            escrever_inline(cels[i].paragraphs[0], valor)
            cels[i].paragraphs[0].runs and setattr(
                cels[i].paragraphs[0].runs[0].font, "size", Pt(10)
            )

    d.add_paragraph()


def adicionar_codigo(d, linhas):
    for linha in linhas:
        p = d.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(linha if linha else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    d.add_paragraph()


def converter(markdown: str) -> "docx.document.Document":
    d = docx.Document()
    configurar_estilos(d)

    linhas = markdown.split("\n")
    i = 0
    primeiro_titulo = True

    while i < len(linhas):
        linha = linhas[i]

        # Bloco de código
        if linha.startswith("```"):
            bloco, i = [], i + 1
            while i < len(linhas) and not linhas[i].startswith("```"):
                bloco.append(linhas[i])
                i += 1
            adicionar_codigo(d, bloco)
            i += 1
            continue

        # Tabela
        if linha.startswith("|") and i + 1 < len(linhas) and set(linhas[i + 1].strip()) <= set("|-: "):
            bloco = []
            while i < len(linhas) and linhas[i].startswith("|"):
                bloco.append(linhas[i])
                i += 1
            adicionar_tabela(d, bloco)
            continue

        # Títulos
        if linha.startswith("#"):
            nivel = len(linha) - len(linha.lstrip("#"))
            texto = linha.lstrip("#").strip()
            if nivel == 1 and primeiro_titulo:
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(texto)
                run.bold = True
                run.font.size = Pt(22)
                run.font.color.rgb = PRIMARIA
                primeiro_titulo = False
            else:
                d.add_heading(texto, level=min(nivel, 3))
            i += 1
            continue

        # Separador
        if linha.strip() == "---":
            i += 1
            continue

        # Lista
        if re.match(r"^\s*[-*]\s+", linha):
            texto = re.sub(r"^\s*[-*]\s+", "", linha)
            recuo = (len(linha) - len(linha.lstrip())) // 2
            p = d.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.6 + 0.5 * recuo)
            escrever_inline(p, texto)
            i += 1
            continue

        if re.match(r"^\s*\d+\.\s+", linha):
            texto = re.sub(r"^\s*\d+\.\s+", "", linha)
            p = d.add_paragraph(style="List Number")
            escrever_inline(p, texto)
            i += 1
            continue

        # Citação
        if linha.startswith(">"):
            p = d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            escrever_inline(p, linha.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        # Parágrafo comum: junta linhas até a próxima em branco
        if linha.strip():
            bloco = []
            while i < len(linhas) and linhas[i].strip() and not linhas[i].startswith(("#", "|", "```", ">", "-", "*")):
                bloco.append(linhas[i].strip())
                i += 1
            escrever_inline(d.add_paragraph(), " ".join(bloco))
            continue

        i += 1

    return d


def main():
    if not FONTE.exists():
        sys.exit(f"Fonte não encontrada: {FONTE}")

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    documento = converter(FONTE.read_text(encoding="utf-8"))
    documento.save(DESTINO)
    print(f"Gerado: {DESTINO}")


if __name__ == "__main__":
    main()
